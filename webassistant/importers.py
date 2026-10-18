import re
from pathlib import Path
from typing import Callable

from .models import BankImport, Candidate, IdentityImport, Question


def load_identities(path: str) -> IdentityImport:
    from openpyxl import load_workbook

    if Path(path).suffix.lower() != ".xlsx":
        raise ValueError("请选择 .xlsx 身份文件")
    wb = load_workbook(path, read_only=True, data_only=True)
    required = ("所在的区县", "姓名", "身份证号")
    try:
        for sheet in wb:
            rows = sheet.iter_rows()
            header = next(rows, ())
            names = [str(c.value or "").strip() for c in header]
            if not all(n in names for n in required):
                continue
            indices = [names.index(n) for n in required]
            candidates, issues, rejected = [], [], []
            for row_number, cells in enumerate(rows, 2):
                chosen = [cells[i].value if i < len(cells) else None for i in indices]
                if all(v is None or str(v).strip() == "" for v in chosen):
                    continue
                district, name, raw_id = chosen
                def reject(reason):
                    issues.append(f"第 {row_number} 行：{reason}")
                    values = ["" if v is None else str(v).strip() for v in chosen]
                    rejected.append((Candidate(row_number, *values), reason))
                if any(v is None or not str(v).strip() for v in chosen):
                    reject("区县、姓名、身份证号均为必填")
                    continue
                if not isinstance(raw_id, str):
                    # Excel keeps only 15 significant decimal digits. Never guess lost ID digits.
                    if isinstance(raw_id, bool) or not isinstance(raw_id, (int, float)) or raw_id != int(raw_id) or len(str(int(raw_id))) != 15:
                        reject("身份证号为数值或格式异常，请从原始资料重新录入为文本；无法恢复已丢失的精度")
                        continue
                    raw_id = str(int(raw_id))
                id_number = raw_id.strip().upper()
                if not re.fullmatch(r"(?:[0-9]{15}|[0-9]{17}[0-9X])", id_number):
                    reject("身份证号应为 15 位或 18 位，末位可为 X")
                    continue
                candidates.append(Candidate(row_number, str(district).strip(), str(name).strip(), id_number))
            return IdentityImport(candidates, issues, sheet.title, rejected)
        raise ValueError("找不到包含「所在的区县、姓名、身份证号」三列表头的工作表")
    finally:
        wb.close()


def parse_question_lines(lines: list[tuple[int, str]], page_count: int) -> BankImport:
    questions, issues = [], []
    current: list[tuple[int, str]] = []
    number = None
    seen: set[int] = set()

    def finish() -> None:
        if number is None:
            return
        pages = tuple(sorted({p for p, _ in current}))
        stem, options, answers = [], {}, ()
        active_option = None
        problems = []
        ended = False
        for _, line in current:
            answer_match = re.fullmatch(r"正确答案\s*[：:]\s*([A-Z]+)\s*", line)
            if answer_match:
                if answers:
                    problems.append("出现多个答案标记")
                answers = tuple(answer_match[1])
                ended = True
                continue
            if ended:
                if line:
                    problems.append("答案后存在未识别内容")
                continue
            option_match = re.match(r"^([A-Z])[.．、]\s*(.*)", line)
            if option_match:
                active_option = option_match[1]
                if active_option in options:
                    problems.append("选项编号重复")
                options[active_option] = option_match[2]
            elif active_option:
                options[active_option] += "\n" + line
            else:
                stem.append(line)
        if not answers:
            problems.append("缺少正确答案")
        if not options or not set(answers).issubset(options):
            problems.append("答案与选项不匹配")
        if not any(stem):
            problems.append("题干为空")
        if number in seen:
            problems.append("题号重复")
        seen.add(number)
        if problems:
            issues.append(f"第 {number} 题（页 {','.join(map(str, pages))}）：{'；'.join(problems)}")
            return
        values = [re.sub(r"\s+", "", v) for v in options.values()]
        kind = "判断" if values == ["正确", "错误"] else "多选" if len(answers) > 1 else "单选"
        questions.append(Question(number, kind, "\n".join(stem), options, answers, pages))

    for page, raw_line in lines:
        line = raw_line.strip()
        if not line or re.fullmatch(r"\d{4}\s*年度法治理论知识考试复习资料", line):
            continue
        start = re.match(r"^(\d+)[、．.]\s*(.*)", line)
        if start:
            finish()
            number = int(start[1])
            current = [(page, start[2])]
        elif number is not None:
            current.append((page, line))
    finish()
    if not questions and not issues:
        issues.append("未发现可识别题目；本版需要包含文字、题号、选项及「正确答案」的 PDF")
    return BankImport(questions, issues, page_count)


def load_bank(path: str, progress: Callable[[str], None] = lambda _: None) -> BankImport:
    import pdfplumber

    if Path(path).suffix.lower() != ".pdf":
        raise ValueError("请选择 PDF 题库")
    lines = []
    with pdfplumber.open(path) as pdf:
        count = len(pdf.pages)
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            page_lines = text.splitlines()
            if page_lines and page_lines[0].strip() == str(i):
                page_lines.pop(0)
            if page_lines and page_lines[-1].strip() == str(i):
                page_lines.pop()
            lines.extend((i, line) for line in page_lines)
            if i % 20 == 0 or i == count:
                progress(f"正在解析题库：{i}/{count} 页")
            page.flush_cache()
    return parse_question_lines(lines, count)
