"""Bounded concurrent verification with durable, resumable identity statistics."""
import asyncio
import json
import os
import threading
import uuid
from dataclasses import asdict, dataclass, replace
from datetime import datetime
from pathlib import Path

from .models import Candidate
from .privacy import redact


LABELS = {"pending": "未开始", "running": "正在登录", "completed": "登录验证完成",
          "failed": "登录失败", "stopped": "已中止，结果未确认"}


def now():
    return datetime.now().isoformat(timespec="seconds")


@dataclass(frozen=True)
class PersonResult:
    index: int
    candidate: Candidate
    status: str = "pending"
    reason: str = ""
    started: str = ""
    ended: str = ""
    stage: str = "登录验证"
    score: int | None = None
    correct: int | None = None
    total: int | None = None
    maximum: int | None = None
    elapsed_seconds: float | None = None
    planned_seconds: float | None = None
    target_score: int | None = None


@dataclass(frozen=True)
class PersonOutcome:
    stage: str
    reason: str
    score: int | None = None
    correct: int | None = None
    total: int | None = None
    maximum: int | None = None
    elapsed_seconds: float | None = None
    planned_seconds: float | None = None
    target_score: int | None = None


def status_text(record):
    if record.status == "completed" and record.score is not None:
        return "本地模拟测试完成"
    if record.status == "running" and record.stage == "本地模拟测试":
        return "正在模拟作答"
    if record.status == "failed" and record.stage == "本地模拟测试":
        return "本地模拟测试失败"
    return LABELS[record.status]


class StatisticsError(RuntimeError):
    pass


class BatchStore:
    def __init__(self, root, candidates, rejected=(), stage="登录验证"):
        self.stage = stage
        self.practice_settings = None
        self.execution_settings = None
        self.folder = Path(root).resolve() / f"登录批次_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}"
        pairs = [(c, "") for c in candidates] + list(rejected)
        pairs.sort(key=lambda pair: pair[0].row)
        self.records = [PersonResult(i, c, "failed" if reason else "pending", reason,
                                     ended=now() if reason else "") for i, (c, reason) in enumerate(pairs)]

    def update(self, index, **changes):
        self.records[index] = replace(self.records[index], **changes)
        self.persist()
        return self.records[index]

    def persist(self):
        try:
            self.folder.mkdir(parents=True, exist_ok=True)
            target = self.folder / "人员状态.json"
            temporary = target.with_suffix(".json.tmp")
            payload = {"version": 5, "stage": self.stage, "updated": now(), "practice_settings": self.practice_settings,
                       "execution_settings": self.execution_settings,
                       "records": [asdict(r) for r in self.records]}
            with temporary.open("w", encoding="utf-8") as stream:
                json.dump(payload, stream, ensure_ascii=False, indent=2)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, target)
        except OSError as error:
            raise StatisticsError("人员状态保存失败，批量任务已停止；请检查结果目录的权限与磁盘空间") from error
        self.export()

    def export(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        try:
            for completed, filename in ((True, "已完成信息统计.xlsx"), (False, "未完成信息统计.xlsx")):
                wb = Workbook()
                ws = wb.active
                ws.title = "人员测试统计"
                ws.append(["所在的区县", "姓名", "身份证号", "源文件行号", "状态", "处理阶段", "原因", "开始时间", "结束时间", "本地模拟得分", "满分", "正确题数", "总题数", "实际用时（分钟）", "计划用时（分钟；0为最快）", "目标分数"])
                for record in self.records:
                    if (record.status == "completed") != completed:
                        continue
                    c = record.candidate
                    values = [c.district, c.name, c.id_number, c.row, status_text(record), record.stage,
                              record.reason, record.started, record.ended, record.score, record.maximum, record.correct, record.total,
                              round(record.elapsed_seconds / 60, 3) if record.elapsed_seconds is not None else None,
                              round(record.planned_seconds / 60, 3) if record.planned_seconds is not None else None, record.target_score]
                    ws.append(values)
                    for cell, value in zip(ws[ws.max_row], values):
                        if isinstance(value, str):
                            # Keep IDs exact; formula-like input is always literal text.
                            cell.data_type = "s"
                            cell.number_format = "@"
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="2464AC")
                for column, width in zip("ABCDEFGHIJKLMNOP", (20, 18, 26, 14, 25, 20, 65, 24, 24, 18, 12, 14, 14, 22, 30, 14)):
                    ws.column_dimensions[column].width = width
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                target = self.folder / filename
                temporary = target.with_suffix(".tmp.xlsx")
                try:
                    wb.save(temporary)
                    os.replace(temporary, target)
                finally:
                    wb.close()
        except (OSError, ValueError) as error:
            raise StatisticsError("Excel 统计导出失败，批量任务已停止；人员状态已保存在 JSON 中。请关闭占用的统计表，点击「重新导出结果」") from error


class BatchControl:
    def __init__(self):
        self.paused = threading.Event()
        self.after_current = threading.Event()
        self.immediate = threading.Event()


class BatchRunner:
    def __init__(self, store, control, changed=lambda _: None):
        self.store, self.control, self.changed = store, control, changed

    def update(self, index, **changes):
        try:
            return self.store.update(index, **changes)
        except StatisticsError:
            self.control.immediate.set()
            if changes.get("status") == "running":
                # No browser action has started yet. Do not leave a phantom active person.
                self.store.records[index] = replace(self.store.records[index], status="failed",
                    reason="统计写入失败，未开始本人员的浏览器操作", ended=now())
                try:
                    self.store.persist()
                except StatisticsError:
                    pass  # Retain the in-memory result so Re-export can recover it.
            raise
        finally:
            # The UI must also show the last known outcome when export is blocked.
            self.changed(self.store.records[index])

    async def run(self, perform, concurrency=1):
        from collections import deque
        from .models import ExecutionOptions
        ExecutionOptions(concurrency).validate()
        self.store.persist()
        pending = deque(r.index for r in self.store.records if r.status == "pending")

        async def worker():
            while pending:
                while self.control.paused.is_set() and not (self.control.immediate.is_set() or self.control.after_current.is_set()):
                    await asyncio.sleep(0.05)
                if not pending or self.control.immediate.is_set() or self.control.after_current.is_set():
                    return
                index = pending.popleft()
                candidate = self.store.records[index].candidate
                self.update(index, status="running", started=now())
                try:
                    outcome = await perform(candidate)
                except asyncio.CancelledError:
                    self.update(index, status="stopped", reason="用户直接停止或批次中断，当前人员结果未确认", ended=now())
                    raise
                except StatisticsError:
                    try:
                        self.update(index, status="failed", reason="统计写入失败，当前人员未完成全部步骤", ended=now())
                    except StatisticsError:
                        pass
                    raise
                except Exception as error:
                    from playwright.async_api import Error as PlaywrightError
                    message = "页面超时、浏览器已关闭或浏览器操作失败" if isinstance(error, PlaywrightError) else str(error)
                    self.update(index, status="failed", reason=redact(message, (candidate.name, candidate.id_number)), ended=now())
                else:
                    changes = asdict(outcome) if isinstance(outcome, PersonOutcome) else {"reason": "已确认考生主页及个人信息标志"}
                    self.update(index, status="completed", ended=now(), **changes)

        tasks = [asyncio.create_task(worker()) for _ in range(min(concurrency, len(pending)))]
        try:
            await asyncio.gather(*tasks)
        except BaseException:
            for task in tasks:
                task.cancel()
            await asyncio.gather(*tasks, return_exceptions=True)
            for record in list(self.store.records):
                if record.status == "running":
                    self.store.records[record.index] = replace(record, status="stopped", reason="批次中断，当前人员结果未确认", ended=now())
                    self.changed(self.store.records[record.index])
            try:
                self.store.persist()
            except StatisticsError:
                pass
            raise
        return self.counts()

    def counts(self):
        return {status: sum(r.status == status for r in self.store.records) for status in LABELS}
