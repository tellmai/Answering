from collections import Counter
from pathlib import Path

import pytest
from openpyxl import Workbook

from webassistant.importers import load_bank, load_identities, parse_question_lines
from webassistant.models import SiteConfig
from webassistant.privacy import redact


ROOT = Path(__file__).resolve().parents[1]


def test_identity_validation_and_first_matching_sheet(tmp_path):
    wb = Workbook()
    sheet = wb.create_sheet("身份")
    sheet.append(["所在的区县", "姓名", "身份证号"])
    sheet.append(["黔江区", "演示甲", "11010119900101001x"])
    sheet.append([None, None, None])
    sheet.append(["黔江区", "演示乙", 110101199001010010])
    sheet.append(["黔江区", "演示丙", "bad-id"])
    sheet.append(["", "演示丁", "110101199001010010"])
    path = tmp_path / "identities.xlsx"
    wb.save(path)
    result = load_identities(str(path))
    assert result.sheet == "身份"
    assert len(result.candidates) == 1
    assert len(result.issues) == 3
    assert len(result.rejected) == 3
    assert [c.row for c, _ in result.rejected] == [4, 5, 6]
    assert result.rejected[1][0].id_number == "bad-id"
    assert result.candidates[0].id_number.endswith("X")
    assert "19900101" not in result.candidates[0].masked_id
    assert "演示甲" not in repr(result.candidates[0])


def test_missing_headers(tmp_path):
    wb = Workbook()
    wb.active.append(["姓名"])
    path = tmp_path / "missing.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="三列表头"):
        load_identities(str(path))


def test_cross_page_options_and_malformed_answers():
    lines = [(1, "1、跨页题干"), (1, "A. 第一项"), (2, "继续第一项"),
             (2, "B. 第二项"), (2, "正确答案：AB"), (2, "2、异常题"),
             (2, "A. 只有 A"), (2, "正确答案：C")]
    result = parse_question_lines(lines, 2)
    assert len(result.questions) == 1 and len(result.issues) == 1
    question = result.questions[0]
    assert question.pages == (1, 2)
    assert question.answers == ("A", "B")
    assert question.kind == "多选"
    assert "继续" in question.options["A"]


def test_numeric_option_continuation_is_preserved():
    result = parse_question_lines([(1, "1、年份"), (1, "A. 数字"), (1, "2025"),
        (1, "B. 其他"), (1, "正确答案：A")], 1)
    assert result.questions[0].options["A"].endswith("2025")


def test_private_error_redaction():
    raw = "身份证 110101199001010010；姓名 演示考生\nCall log: fill(secret)"
    safe = redact(raw, ("演示考生",))
    assert "110101199001010010" not in safe
    assert "演示考生" not in safe
    assert "fill(secret)" not in safe


@pytest.mark.parametrize("url", ["http://ks.cqsdx.cn/exam/", "https://example.com/exam/", "https://ks.cqsdx.cn.evil.example/exam/", "https://ks.cqsdx.cn/exam/?id=123"])
def test_wrong_destination_rejected(url):
    with pytest.raises(ValueError):
        SiteConfig(url).validate()


def test_provided_identity():
    path = ROOT / "filedata/身份信息.xlsx"
    if not path.exists():
        pytest.skip("private source fixture is not distributed")
    result = load_identities(str(path))
    assert result.candidates
    assert not result.issues
    assert all(c.district and c.name and isinstance(c.id_number, str) for c in result.candidates)


def test_provided_bank():
    path = ROOT / "filedata/题库.pdf"
    if not path.exists():
        pytest.skip("private source fixture is not distributed")
    result = load_bank(str(path))
    assert result.page_count == 276
    assert not result.issues, result.issues[:10]
    assert [q.number for q in result.questions] == list(range(1, 1022))
    assert Counter(q.kind for q in result.questions) == {"判断": 324, "单选": 368, "多选": 329}
    from webassistant.mock_exam import validate_bank, sample_questions
    knowledge = validate_bank(result.questions)
    for question in result.questions:
        assert knowledge.answer(question.kind, question.stem, question.options) == sorted(question.answers)
    assert len(sample_questions(result.questions)) == 100
