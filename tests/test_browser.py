import asyncio
import base64
import os
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse

import pytest

from webassistant.browser import LoginFailure, SiteSession, find_chrome
from webassistant.models import Candidate, SiteConfig, PracticeOptions, ExecutionOptions


# This is a local fixture, not a real challenge. No private data or real server is used.
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jC1sAAAAASUVORK5CYII=")
ENTRY = '''<div id="districtGrid"><article class="district-card" role="button" aria-label="进入黔江区考试入口" onclick="location.href='/exam/qianjiang'"><div class="district-name">黔江区</div></article></div>'''
LOGIN = '''<div id="roleSwitch"><button data-role="admin" class="active">工作人员登录</button><button data-role="student" onclick="this.className='active';document.querySelector('#username').placeholder='请输入姓名';document.querySelector('#password').placeholder='请输入身份证号码'">考生登录</button></div>
<input id="username"><input id="password"><input id="captchaInput"><div id="captchaImage" onclick="refresh()"><img id="captcha" src="/fixture.png?n=0"></div>
<div id="loginNotice" style="display:none"></div><div id="pageLoading" style="display:none"></div>
<button id="loginButton" onclick="login()">进入考试</button>
<script>let n=0;function refresh(){document.querySelector('#captcha').src='/fixture.png?n='+(++n)}
async function login(){document.querySelector('#loginNotice').style.display='none';const r=await fetch('/mock-submit',{method:'POST',body:document.querySelector('#captchaInput').value});const t=await r.text();if(t==='ok'){location.href='/exam/user/home'}else{refresh();const e=document.querySelector('#loginNotice');e.style.display='block';e.textContent=t;}}
</script>'''
HOME = '<h1>您好，演示考生</h1><a>个人资料</a>'


@pytest.fixture
def local_site():
    state = {"posts": 0, "mode": "normal"}
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            path = urlparse(self.path).path
            if path == "/delay.js":
                time.sleep(0.4)
                body, mime = b"/* delayed frontend fixture */", "text/javascript"
            elif path == "/fixture.png":
                body, mime = PNG, "image/png"
            else:
                html = ENTRY if path == "/exam/" else HOME if path == "/exam/user/home" else LOGIN
                if state.get("delayed") and path == "/exam/qianjiang":
                    inline = "onclick=\"this.className='active';document.querySelector('#username').placeholder='请输入姓名';document.querySelector('#password').placeholder='请输入身份证号码'\""
                    html = html.replace(inline, "")
                    handler = "document.querySelector('[data-role=student]').onclick=function(){this.className='active';document.querySelector('#username').placeholder='请输入姓名';document.querySelector('#password').placeholder='请输入身份证号码';};"
                    html = html.replace("<script>", '<script src="/delay.js"></script><script>' + handler)
                if state.get("delayed_captcha") and path == "/exam/qianjiang":
                    html = html.replace("function refresh(){", "async function refresh(){await new Promise(r=>setTimeout(r,350));")
                body, mime = html.encode(), "text/html; charset=utf-8"
            self.send_response(200)
            self.send_header("Content-Type", mime)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        def do_POST(self):
            state["posts"] += 1
            code = self.rfile.read(int(self.headers.get("Content-Length", 0))).decode()
            text = "姓名或身份证号不匹配" if state["mode"] == "identity-error" else "ok" if code == "TEST" else "验证码错误"
            body = text.encode()
            self.send_response(200)
            self.send_header("Content-Type", "text/plain; charset=utf-8")
            self.end_headers()
            self.wfile.write(body)
        def log_message(self, *_):
            pass
    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    yield f"http://127.0.0.1:{server.server_port}/exam/", state
    server.shutdown()
    server.server_close()
    thread.join()


class FakeOCR:
    def __init__(self, code):
        self.code = code
    async def recognize(self, _):
        return self.code


def run_scenario(url, scenario):
    async def run():
        from playwright.async_api import async_playwright
        executable = os.environ.get("TEST_CHROME_PATH") or find_chrome()
        if not executable:
            pytest.skip("Chrome missing; set TEST_CHROME_PATH")
        async with async_playwright() as p:
            session = SiteSession(p, SiteConfig(url, executable, 5000, True))
            # Use headless Chrome for isolated fixture tests, same methods as production.
            real_open = session.open
            async def headless_open():
                await real_open(headless=True)
            session.open = headless_open
            try:
                await scenario(session)
            finally:
                await session.close()
    asyncio.run(run())


CANDIDATE = Candidate(2, "黔江区", "演示考生", "110101199001010010")


def test_success_keeps_browser_and_never_opens_exam(local_site):
    url, state = local_site
    async def scenario(session):
        async def unexpected(_):
            raise AssertionError("manual input should not be requested")
        message = await session.login(CANDIDATE, FakeOCR("TEST"), unexpected)
        assert "登录成功" in message
        assert session.browser.is_connected()
        assert urlparse(session.page.url).path == "/exam/user/home"
    run_scenario(url, scenario)
    assert state["posts"] == 1


def test_three_failures_then_manual_input(local_site):
    url, state = local_site
    async def scenario(session):
        async def manual(image):
            assert image.startswith(b"\x89PNG")
            assert state["posts"] == 3
            return "TEST"
        assert "登录成功" in await session.login(CANDIDATE, FakeOCR("wrong"), manual)
    run_scenario(url, scenario)
    assert state["posts"] == 4


def test_batch_captcha_limit_records_failure_then_next_person(local_site, tmp_path):
    from dataclasses import replace
    from webassistant.batch import BatchStore, BatchControl, BatchRunner
    url, state = local_site
    async def scenario(session):
        store = BatchStore(tmp_path, [CANDIDATE, replace(CANDIDATE, row=3)])
        async def perform(candidate):
            await session.close()
            await session.login(candidate, FakeOCR("wrong" if candidate.row == 2 else "TEST"))
        counts = await BatchRunner(store, BatchControl()).run(perform)
        assert counts["failed"] == 1 and counts["completed"] == 1
        assert "三次" in store.records[0].reason
    run_scenario(url, scenario)
    assert state["posts"] == 4


def test_captcha_retry_waits_for_new_image(local_site):
    url, state = local_site
    state["delayed_captcha"] = True
    async def scenario(session):
        class InspectOCR:
            calls = 0
            async def recognize(self, _):
                self.calls += 1
                src = await session.page.locator("#captcha").get_attribute("src")
                if self.calls == 1:
                    assert src.endswith("n=0")
                    return "wrong"
                assert src.endswith("n=1")
                return "TEST"
        assert "登录成功" in await session.login(CANDIDATE, InspectOCR())
    run_scenario(url, scenario)
    assert state["posts"] == 2


def test_worker_login_to_local_exam_scores_two_people_and_closes_pages(local_site, sample_bank, tmp_path, monkeypatch):
    from dataclasses import replace
    from PySide6.QtWidgets import QApplication
    from PySide6.QtCore import QTimer
    from openpyxl import load_workbook
    from webassistant.controller import BrowserWorker
    url, state = local_site
    class HeadlessSession(SiteSession):
        async def open(self, headless=True):
            await super().open(headless=True)
    monkeypatch.setattr("webassistant.controller.SiteSession", HeadlessSession)
    app = QApplication.instance() or QApplication([])
    worker = BrowserWorker()
    ready, finished, errors, scores = [], [], [], []
    worker.ready.connect(lambda: ready.append(True))
    worker.job_finished.connect(lambda: finished.append(True))
    worker.failed.connect(errors.append)
    worker.score_ready.connect(lambda candidate, result: scores.append((candidate.row, result.score)))
    live, active_counts, ticks = {}, [], []
    def updated(record):
        live[record.index] = record.status
        active_counts.append(sum(value == "running" for value in live.values()))
    worker.person_updated.connect(updated)
    heartbeat = QTimer()
    heartbeat.setInterval(10)
    heartbeat.timeout.connect(lambda: ticks.append(True))
    heartbeat.start()
    async def recognize(_):
        return "TEST"
    monkeypatch.setattr(worker.ocr, "recognize", recognize)
    def wait(condition, timeout=45):
        deadline = time.monotonic() + timeout
        while not condition() and time.monotonic() < deadline:
            app.processEvents()
            time.sleep(0.01)
        assert condition()
    worker.start()
    try:
        wait(lambda: ready)
        worker.start_job("login", SiteConfig(url, find_chrome(), 10000, True),
                         [CANDIDATE, replace(CANDIDATE, row=3)], str(tmp_path), (), sample_bank, PracticeOptions(fastest=True, score_min=99, score_max=100), ExecutionOptions(concurrency=2, background=True))
        wait(lambda: finished)
        assert not errors
        assert max(active_counts) == 2 and len(ticks) > 5
        assert sorted(row for row, _ in scores) == [2, 3] and {score for _, score in scores} == {99, 100}
        assert all(r.status == "completed" and r.stage == "本地模拟测试" and 99 <= r.score <= 100 and r.total == 100 for r in worker.last_store.records)
        assert all(s.page.is_closed() for s in worker.sessions)
        assert state["posts"] == 2
        wb = load_workbook(worker.last_store.folder / "已完成信息统计.xlsx")
        assert sorted(r[9:13] for r in wb.active.iter_rows(min_row=2, values_only=True)) == [(99, 100, 99, 100), (100, 100, 100, 100)]
        wb.close()
    finally:
        heartbeat.stop()
        worker.shutdown()
        wait(lambda: not worker.isRunning())


def test_shared_browser_contexts_isolate_cookies_and_close_independently(local_site):
    from playwright.async_api import async_playwright
    url, state = local_site
    async def run():
        async with async_playwright() as playwright:
            browser = await playwright.chromium.launch(executable_path=find_chrome(), headless=True)
            sessions = [SiteSession(playwright, SiteConfig(url, find_chrome(), 10000, True), browser=browser) for _ in range(2)]
            try:
                await asyncio.gather(*(session.login(CANDIDATE, FakeOCR("TEST")) for session in sessions))
                assert sessions[0].context != sessions[1].context
                for index, session in enumerate(sessions):
                    await session.context.add_cookies([{"name": "identity", "value": str(index), "url": url}])
                for index, session in enumerate(sessions):
                    cookies = await session.context.cookies()
                    assert next(c["value"] for c in cookies if c["name"] == "identity") == str(index)
                await sessions[0].close()
                assert browser.is_connected() and not sessions[1].page.is_closed()
            finally:
                await asyncio.gather(*(session.close() for session in sessions))
                await browser.close()
    asyncio.run(run())
    assert state["posts"] == 2


def test_identity_error_is_not_retried(local_site):
    url, state = local_site
    state["mode"] = "identity-error"
    async def scenario(session):
        with pytest.raises(LoginFailure, match="不匹配"):
            await session.login(CANDIDATE, FakeOCR("TEST"), None)
    run_scenario(url, scenario)
    assert state["posts"] == 1


def test_stop_while_waiting_manual_prevents_more_posts(local_site):
    url, state = local_site
    async def scenario(session):
        waiting = asyncio.Event()
        async def manual(_):
            waiting.set()
            await asyncio.Future()
        task = asyncio.create_task(session.login(CANDIDATE, FakeOCR("wrong"), manual))
        await asyncio.wait_for(waiting.wait(), 10)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert session.browser.is_connected()
        await asyncio.sleep(0.2)
        assert state["posts"] == 3
    run_scenario(url, scenario)


def test_missing_district_and_closed_browser(local_site):
    url, state = local_site
    async def scenario(session):
        await session.open()
        with pytest.raises(LoginFailure, match="区县"):
            await session.choose_district("不存在区")
        await session.page.close()
        from playwright.async_api import Error
        with pytest.raises(Error):
            await session.choose_district("黔江区")
    run_scenario(url, scenario)
    assert state["posts"] == 0


def test_browser_closed_during_manual_wait_ends_job(local_site):
    url, state = local_site
    async def scenario(session):
        waiting = asyncio.Event()
        async def manual(_):
            waiting.set()
            await asyncio.Future()
        task = asyncio.create_task(session.login(CANDIDATE, FakeOCR("wrong"), manual))
        await asyncio.wait_for(waiting.wait(), 10)
        await session.page.close()
        with pytest.raises(LoginFailure, match="已关闭"):
            await asyncio.wait_for(task, 5)
    run_scenario(url, scenario)
    assert state["posts"] == 3


def test_waits_for_delayed_frontend_before_switching_role(local_site):
    url, state = local_site
    state["delayed"] = True
    async def scenario(session):
        assert "登录成功" in await session.login(CANDIDATE, FakeOCR("TEST"), None)
    run_scenario(url, scenario)
    assert state["posts"] == 1


def test_stop_during_timed_exam_keeps_real_page_and_server(local_site, sample_bank, tmp_path, monkeypatch):
    import json
    from urllib.request import urlopen
    from PySide6.QtWidgets import QApplication
    from webassistant.controller import BrowserWorker
    url, state = local_site
    class HeadlessSession(SiteSession):
        async def open(self, headless=True):
            await super().open(headless=True)
    monkeypatch.setattr("webassistant.controller.SiteSession", HeadlessSession)
    app = QApplication.instance() or QApplication([])
    worker = BrowserWorker()
    ready, finished, steps = [], [], []
    worker.ready.connect(lambda: ready.append(True))
    worker.job_finished.connect(lambda: finished.append(True))
    worker.progress.connect(steps.append)
    async def recognize(_):
        return "TEST"
    monkeypatch.setattr(worker.ocr, "recognize", recognize)
    def wait(condition):
        deadline = time.monotonic() + 20
        while not condition() and time.monotonic() < deadline:
            app.processEvents()
            time.sleep(0.01)
        assert condition()
    worker.start()
    page = None
    try:
        wait(lambda: ready)
        worker.start_job("login", SiteConfig(url, find_chrome(), 10000, True), [CANDIDATE], str(tmp_path), (), sample_bank,
            PracticeOptions(minutes_min=1, minutes_max=1))
        wait(lambda: any("本地模拟作答：5/100" in step for step in steps))
        worker.stop_job()
        wait(lambda: finished)
        assert worker.last_store.records[0].status == "stopped"
        assert len(worker.retained_sessions) == len(worker.retained_servers) == 1
        session, server = worker.retained_sessions[0], worker.retained_servers[0]
        page = session.page
        assert not page.is_closed() and session.browser.is_connected()
        paper = next(iter(server.papers.values()))
        with urlopen(server.url_for(paper) + "data") as response:
            assert len(json.loads(response.read())["questions"]) == 100
        assert server.result_for(paper) is None and state["posts"] == 1
        time.sleep(0.2)
        assert not page.is_closed() and server.result_for(paper) is None
    finally:
        worker.shutdown()
        wait(lambda: not worker.isRunning())
        if page:
            assert page.is_closed()
