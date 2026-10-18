import asyncio
import json

import pytest
from openpyxl import load_workbook

from webassistant.batch import BatchControl, BatchRunner, BatchStore, StatisticsError
from webassistant.browser import CaptchaFailure, LoginFailure
from webassistant.models import Candidate


def people(count=3):
    return [Candidate(i + 2, "黔江区", f"演示人员{i}", f"1101011990010100{i:02}") for i in range(count)]


def test_failures_continue_and_exports_preserve_all_people(tmp_path):
    records = people()
    invalid = Candidate(8, "", "=HYPERLINK(\"demo\")", "bad-id")
    store = BatchStore(tmp_path, records, [(invalid, "缺少区县")])
    visited = []
    async def perform(c):
        visited.append(c.row)
        if c.row == 2:
            raise CaptchaFailure("验证码连续三次失败")
        if c.row == 3:
            raise LoginFailure("姓名或身份证号不匹配 " + c.id_number)
    counts = asyncio.run(BatchRunner(store, BatchControl()).run(perform))
    assert visited == [2, 3, 4]
    assert counts["completed"] == 1 and counts["failed"] == 3
    assert records[1].id_number not in store.records[1].reason
    wb = load_workbook(store.folder / "已完成信息统计.xlsx")
    assert wb.active["C2"].value == records[2].id_number
    assert wb.active["C2"].data_type == "s" and wb.active["C2"].number_format == "@"
    wb.close()
    wb = load_workbook(store.folder / "未完成信息统计.xlsx")
    assert wb.active.max_row == 4
    assert wb.active["B4"].value == invalid.name and wb.active["B4"].data_type == "s"
    wb.close()


def test_graceful_stop_finishes_current_and_keeps_remaining(tmp_path):
    store, control = BatchStore(tmp_path, people()), BatchControl()
    entered, release = asyncio.Event(), asyncio.Event()
    visited = []
    async def perform(c):
        visited.append(c.row)
        entered.set()
        await release.wait()
    async def run():
        task = asyncio.create_task(BatchRunner(store, control).run(perform))
        await entered.wait()
        control.after_current.set()
        assert not task.done()
        release.set()
        return await task
    counts = asyncio.run(run())
    assert visited == [2]
    assert counts["completed"] == 1 and counts["pending"] == 2
    wb = load_workbook(store.folder / "未完成信息统计.xlsx")
    assert [r[4] for r in wb.active.iter_rows(min_row=2, values_only=True)] == ["未开始", "未开始"]
    wb.close()


def test_immediate_stop_persists_interrupted_and_pending(tmp_path):
    store = BatchStore(tmp_path, people())
    entered = asyncio.Event()
    async def perform(_):
        entered.set()
        await asyncio.Future()
    async def run():
        task = asyncio.create_task(BatchRunner(store, BatchControl()).run(perform))
        await entered.wait()
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
    asyncio.run(run())
    data = json.loads((store.folder / "人员状态.json").read_text(encoding="utf-8"))
    assert [r["status"] for r in data["records"]] == ["stopped", "pending", "pending"]


@pytest.mark.parametrize("stop", [False, True])
def test_stop_dialog_gate_never_starts_next_until_choice(tmp_path, stop):
    store, control = BatchStore(tmp_path, people(2)), BatchControl()
    visited = []
    entered, release = asyncio.Event(), asyncio.Event()
    async def perform(c):
        visited.append(c.row)
        if c.row == 2:
            entered.set()
            await release.wait()
    async def run():
        task = asyncio.create_task(BatchRunner(store, control).run(perform))
        await entered.wait()
        control.paused.set()
        release.set()
        await asyncio.sleep(0.15)
        assert visited == [2] and not task.done()
        if stop:
            control.after_current.set()
        control.paused.clear()
        await task
    asyncio.run(run())
    assert visited == ([2] if stop else [2, 3])


def test_stop_before_start_still_exports_queue(tmp_path):
    store, control = BatchStore(tmp_path, people()), BatchControl()
    control.immediate.set()
    async def unexpected(_):
        raise AssertionError("must not start")
    counts = asyncio.run(BatchRunner(store, control).run(unexpected))
    assert counts["pending"] == 3
    assert (store.folder / "未完成信息统计.xlsx").exists()


def test_export_failure_stops_queue_but_json_keeps_latest_outcome(tmp_path, monkeypatch):
    store = BatchStore(tmp_path, people(2))
    original = store.export
    def blocked_export():
        if any(r.status == "completed" for r in store.records):
            raise StatisticsError("模拟统计文件被占用")
        original()
    monkeypatch.setattr(store, "export", blocked_export)
    visited = []
    async def perform(c):
        visited.append(c.row)
    with pytest.raises(StatisticsError):
        asyncio.run(BatchRunner(store, BatchControl()).run(perform))
    assert visited == [2]
    saved = json.loads((store.folder / "人员状态.json").read_text(encoding="utf-8"))
    assert [r["status"] for r in saved["records"]] == ["completed", "pending"]
    monkeypatch.setattr(store, "export", original)
    store.persist()
    assert (store.folder / "已完成信息统计.xlsx").exists()


def test_statistics_failure_before_browser_does_not_leave_running_person(tmp_path, monkeypatch):
    store = BatchStore(tmp_path, people(2))
    original = store.export
    def blocked_export():
        if any(r.status == "running" for r in store.records):
            raise StatisticsError("模拟启动前无法写入统计")
        original()
    monkeypatch.setattr(store, "export", blocked_export)
    async def unexpected(_):
        raise AssertionError("browser must not start")
    with pytest.raises(StatisticsError):
        asyncio.run(BatchRunner(store, BatchControl()).run(unexpected))
    assert [r.status for r in store.records] == ["failed", "pending"]
    assert "未开始" in store.records[0].reason
